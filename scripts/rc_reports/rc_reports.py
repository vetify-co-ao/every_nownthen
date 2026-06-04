# /// script
# requires-python = ">=3.12"
# dependencies = [
#     "requests",
#     "openpyxl",
#     "google-api-python-client",
#     "google-auth",
#     "google-auth-httplib2",
# ]
# ///
import base64
import datetime as dt
import json
import os
import shutil
import sys
from dataclasses import dataclass
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path

import requests
from openpyxl import load_workbook

# --- Configuration Constants ---

SCRIPT_DIR = Path(__file__).resolve().parent
XLSX_TEMPLATE_PATH = SCRIPT_DIR / "RC_Report_Template_2026.xlsx"

# Dashy API
DASHY_API_KEY = os.environ["DASHY_API_KEY"]
DASHY_API_BASE_URL = "https://bi.vetify.co.ao"
DASHY_TIMEOUT_SECONDS = 60

# Period configuration
DEFAULT_PERIOD_START_DATE = f"{dt.date.today().year}-01-01"
PERIOD_START_DATE = os.environ.get(
    "RC_REPORTS_PERIOD_START_DATE", DEFAULT_PERIOD_START_DATE
)

# Output
OUTPUT_DIR = Path(os.environ.get("RC_REPORTS_OUTPUT_DIR", "/tmp"))

# Email / Gmail API
SERVICE_ACCOUNT_KEY_PATH = os.environ["SERVICE_ACCOUNT_KEY_PATH"]
RC_REPORTS_EMAIL_TO = os.environ["RC_REPORTS_EMAIL_TO"]
IMPERSONATED_EMAIL = "comercial@vetify.co.ao"
EMAIL_FROM = "Vetify <comercial@vetify.co.ao>"
REPLY_TO = "encomendas@vetify.co.ao"
EMAIL_SUBJECT_TEMPLATE = "Royal Canin Report P%s (%s a %s)"
GMAIL_API_SCOPES = [
    "https://www.googleapis.com/auth/gmail.send",
    "https://mail.google.com/",
]

# Excel layout
EXCEL_SHEET_NAME = "Sheet1"
EXCEL_FIRST_DATA_ROW = 2
EXCEL_COLUMNS = {
    "sku": "A",
    "product_name": "B",
    "stock_on_hand": "C",
    "sales_usd": "D",
    "sales_kg": "E",
}


@dataclass(frozen=True)
class ReportPeriod:
    number: int
    start: dt.date
    end: dt.date


@dataclass(frozen=True)
class ReportRow:
    sku: str
    product_name: str
    stock_on_hand: float
    sales_usd: float
    sales_kg: float
    units_sold: float
    sales_aoa: float


def parse_date(value: str) -> dt.date:
    try:
        return dt.date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError(f"Invalid date '{value}'. Expected YYYY-MM-DD.") from exc


def calculate_period(period_number: int, base_date: dt.date) -> ReportPeriod:
    if period_number < 1 or period_number > 13:
        raise ValueError("Period number must be between 1 and 13.")

    start = base_date + dt.timedelta(days=(period_number - 1) * 28)
    end = start + dt.timedelta(days=27)
    return ReportPeriod(period_number, start, end)


def to_float(value, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def is_royal_canin_product(product: dict) -> bool:
    return str(product.get("brand", "")).strip().casefold() == "royal canin"


class DashyClient:
    def __init__(self, api_key: str, base_url: str = DASHY_API_BASE_URL, session=None):
        self.api_key = api_key
        self.base_url = base_url.rstrip("/")
        self.session = session or requests.Session()
        self.session.headers.update(
            {"Accept": "application/json", "User-Agent": "every-nownthen-rc-reports/1.0"}
        )

    def _get(self, path: str, params: dict | None = None, include_api_key: bool = False):
        request_params = dict(params or {})
        if include_api_key:
            request_params["api_key"] = self.api_key

        try:
            response = self.session.get(
                f"{self.base_url}{path}",
                params=request_params,
                timeout=DASHY_TIMEOUT_SECONDS,
            )
            response.raise_for_status()
            return response.json()
        except requests.exceptions.RequestException as exc:
            details = getattr(exc.response, "text", "") if getattr(exc, "response", None) else ""
            details = details.replace(self.api_key, "***")
            status = getattr(getattr(exc, "response", None), "status_code", "unknown")
            raise RuntimeError(
                f"Dashy request failed for {path} (status={status}). {details[:500]}"
            ) from exc

    def list_products(self) -> list[dict]:
        # The Dashy OpenAPI documents api_key as a query auth mechanism and notes
        # it creates a temporary Dashy session cookie. Dataset/metric calls then
        # use that cookie and must not repeat api_key as a business parameter.
        payload = self._get("/api/dimensions/product", include_api_key=True)
        return payload.get("rows", [])

    def get_sales_by_sku(
        self,
        start_date: dt.date,
        end_date: dt.date,
        brand: str = "Royal Canin",
        limit: int = 100,
    ) -> dict[str, dict[str, float]]:
        payload = self._get(
            "/api/datasets/client_top_products",
            params={
                "date.from": start_date.isoformat(),
                "date.until": end_date.isoformat(),
                "brand": brand,
                "limit": min(limit, 100),
            },
        )
        skus = payload.get("x", [])
        series = payload.get("series", {})
        revenues = series.get("revenue_aoa", [])
        units = series.get("units_sold", [])

        sales: dict[str, dict[str, float]] = {}
        for index, sku in enumerate(skus):
            sales[str(sku)] = {
                "revenue_aoa": to_float(revenues[index] if index < len(revenues) else 0),
                "units_sold": to_float(units[index] if index < len(units) else 0),
            }
        return sales

    def get_stock_on_date(self, sku: str, snapshot_date: dt.date) -> float:
        snapshot_date_str = snapshot_date.isoformat()
        payload = self._get(
            "/api/datasets/product_stock_daily",
            params={
                "date.from": snapshot_date_str,
                "date.until": snapshot_date_str,
                "product_sku": sku,
            },
        )
        dates = payload.get("x", [])
        stock_values = payload.get("series", {}).get("stock_qty", [])
        if snapshot_date_str not in dates:
            return 0
        index = dates.index(snapshot_date_str)
        if index >= len(stock_values):
            return 0
        return to_float(stock_values[index])

    def get_usd_rate_on_date(self, snapshot_date: dt.date) -> float:
        snapshot_date_str = snapshot_date.isoformat()
        payload = self._get(
            "/api/datasets/currency_rates_vs_units",
            params={
                "date.from": snapshot_date_str,
                "date.until": snapshot_date_str,
            },
        )
        rates = payload.get("series", {}).get("avg_usd_rate", [])
        usd_rate = to_float(rates[-1] if rates else None)
        if usd_rate <= 0:
            raise ValueError(f"No valid USD exchange rate for {snapshot_date_str}.")
        return usd_rate


def build_report_rows(client: DashyClient, period: ReportPeriod) -> list[ReportRow]:
    products = [product for product in client.list_products() if is_royal_canin_product(product)]
    usd_rate = client.get_usd_rate_on_date(period.end)
    sales_by_sku = client.get_sales_by_sku(
        period.start,
        period.end,
        brand="Royal Canin",
        limit=max(len(products), 1),
    )

    rows: list[ReportRow] = []
    for product in products:
        sku = str(product.get("sku") or "").strip()
        product_name = str(product.get("product_name") or "").strip()
        if not sku or not product_name:
            continue

        sales = sales_by_sku.get(sku, {})
        sales_aoa = to_float(sales.get("revenue_aoa"))
        units_sold = to_float(sales.get("units_sold"))
        stock_on_hand = client.get_stock_on_date(sku, period.end)

        if units_sold == 0 and stock_on_hand == 0:
            continue

        rc_weight = to_float(product.get("rc_weight"))
        rows.append(
            ReportRow(
                sku=sku,
                product_name=product_name,
                stock_on_hand=stock_on_hand,
                sales_usd=sales_aoa / usd_rate,
                sales_kg=units_sold * rc_weight,
                units_sold=units_sold,
                sales_aoa=sales_aoa,
            )
        )

    return sorted(rows, key=lambda row: row.product_name.casefold())


def write_report_workbook(
    template_path: Path,
    output_dir: Path,
    period: ReportPeriod,
    rows: list[ReportRow],
) -> Path:
    if not template_path.exists():
        raise FileNotFoundError(f"XLSX template file not found: {template_path}")

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / (
        f"RC_Report_P{period.number:02d}_{period.start.isoformat()}_{period.end.isoformat()}.xlsx"
    )
    shutil.copy(template_path, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook[EXCEL_SHEET_NAME]

    if sheet.max_row >= EXCEL_FIRST_DATA_ROW:
        for row in sheet.iter_rows(
            min_row=EXCEL_FIRST_DATA_ROW,
            max_row=sheet.max_row,
            min_col=1,
            max_col=len(EXCEL_COLUMNS),
        ):
            for cell in row:
                cell.value = None

    for index, row in enumerate(rows, start=EXCEL_FIRST_DATA_ROW):
        sheet[f"{EXCEL_COLUMNS['sku']}{index}"] = row.sku
        sheet[f"{EXCEL_COLUMNS['product_name']}{index}"] = row.product_name
        sheet[f"{EXCEL_COLUMNS['stock_on_hand']}{index}"] = row.stock_on_hand
        sheet[f"{EXCEL_COLUMNS['sales_usd']}{index}"] = row.sales_usd
        sheet[f"{EXCEL_COLUMNS['sales_kg']}{index}"] = row.sales_kg
        sheet[f"{EXCEL_COLUMNS['stock_on_hand']}{index}"].number_format = "#,##0"
        sheet[f"{EXCEL_COLUMNS['sales_usd']}{index}"].number_format = "#,##0.00"
        sheet[f"{EXCEL_COLUMNS['sales_kg']}{index}"].number_format = "#,##0.00"

    workbook.save(output_path)
    return output_path


def parse_email_recipients(value: str) -> list[str]:
    return [email.strip() for email in value.split(",") if email.strip()]


def create_gmail_service():
    if not os.path.exists(SERVICE_ACCOUNT_KEY_PATH):
        raise FileNotFoundError(
            f"Service account key file not found: {SERVICE_ACCOUNT_KEY_PATH}"
        )

    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build

    creds = Credentials.from_service_account_file(
        SERVICE_ACCOUNT_KEY_PATH,
        scopes=GMAIL_API_SCOPES,
        subject=IMPERSONATED_EMAIL,
    )
    return build("gmail", "v1", credentials=creds)


def build_email_message(recipients: list[str], subject: str, html_body: str, attachment_path: Path):
    message = MIMEMultipart()
    message["to"] = ", ".join(recipients)
    message["reply-to"] = REPLY_TO
    message["from"] = EMAIL_FROM
    message["subject"] = subject
    message.attach(MIMEText(html_body, "html"))

    with open(attachment_path, "rb") as attachment_file:
        part = MIMEBase("application", "octet-stream")
        part.set_payload(attachment_file.read())
    encoders.encode_base64(part)
    part.add_header(
        "Content-Disposition",
        f'attachment; filename="{attachment_path.name}"',
    )
    message.attach(part)
    return message


def send_email(gmail_service, recipients: list[str], subject: str, html_body: str, attachment_path: Path):
    if not recipients:
        raise ValueError("At least one email recipient is required.")

    message = build_email_message(recipients, subject, html_body, attachment_path)
    raw_message = base64.urlsafe_b64encode(message.as_bytes()).decode()

    try:
        sent_message = (
            gmail_service.users()
            .messages()
            .send(userId="me", body={"raw": raw_message})
            .execute()
        )
        print(f"Email sent successfully! Message ID: {sent_message['id']}")
    except Exception as exc:
        try:
            from googleapiclient.errors import HttpError
        except Exception:  # pragma: no cover - defensive fallback if import fails
            HttpError = None

        if HttpError and isinstance(exc, HttpError):
            details = exc.resp.get("content", b"{}")
            try:
                print(json.dumps(json.loads(details.decode("utf-8")), indent=2))
            except Exception:
                print(f"Raw Gmail error content: {details}")
        raise


def build_email_body(period: ReportPeriod, row_count: int) -> str:
    return f"""
    <p>Segue em anexo o relatório Royal Canin.</p>
    <ul>
      <li>Período: {period.number}</li>
      <li>Intervalo: {period.start.isoformat()} a {period.end.isoformat()}</li>
      <li>Linhas incluídas: {row_count}</li>
    </ul>
    """


def run(period_number: int) -> Path:
    period = calculate_period(period_number, parse_date(PERIOD_START_DATE))
    recipients = parse_email_recipients(RC_REPORTS_EMAIL_TO)
    if not recipients:
        raise ValueError("RC_REPORTS_EMAIL_TO must contain at least one recipient.")

    print("Starting Royal Canin report script...")
    print(f"Period: P{period.number:02d} ({period.start} to {period.end})")
    print(f"Output directory: {OUTPUT_DIR}")
    print(f"Recipients: {len(recipients)}")

    client = DashyClient(DASHY_API_KEY)
    rows = build_report_rows(client, period)
    print(f"Report rows: {len(rows)}")

    workbook_path = write_report_workbook(XLSX_TEMPLATE_PATH, OUTPUT_DIR, period, rows)
    print(f"Workbook written to: {workbook_path}")

    subject = EMAIL_SUBJECT_TEMPLATE % (
        period.number,
        period.start.isoformat(),
        period.end.isoformat(),
    )
    gmail_service = create_gmail_service()
    send_email(gmail_service, recipients, subject, build_email_body(period, len(rows)), workbook_path)

    print("Royal Canin report script finished successfully.")
    return workbook_path


def print_usage():
    print("Usage: uv run rc_reports.py <period_number>")
    print("  <period_number> must be an integer from 1 to 13.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print_usage()
        sys.exit(2)

    try:
        selected_period = int(sys.argv[1])
        run(selected_period)
    except Exception as exc:
        print(f"Error: {exc}")
        sys.exit(1)
