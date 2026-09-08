import datetime as dt
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from models import Announcement, BulletinData, Cargo, Product
from renderers import (
    build_email_message,
    build_xlsx,
    markdown_to_email_html,
    render_html,
)


def product(**overrides):
    values = {
        "sku": "SKU-1",
        "slug": "product-one",
        "title": "Product One",
        "description": "Useful product",
        "availability": "EM STOCK",
        "net_price": 100.0,
        "vat_rate": 0.14,
        "expiry": dt.date(2026, 12, 31),
        "categories": "Medicamentos",
        "subcategories": "Dermatologia",
        "brand": "Brand",
        "department": "Medicamentos",
        "featured": True,
        "image_url": "https://images.example/product.png",
    }
    values.update(overrides)
    return Product(**values)


def bulletin(**overrides):
    values = {
        "publication_date": dt.date(2026, 9, 7),
        "issue_number": 73,
        "recipients": ["one@example.com"],
        "products": [product()],
        "announcements": [],
    }
    values.update(overrides)
    return BulletinData(**values)


def html_template():
    return """<!doctype html><html><head><title>Boletim Semanal Vetify</title></head><body>
<p>Boletim Semanal · N.º 00 · 00/00/0000</p>
<!-- ══════════ SECÇÃO: INFORMAÇÃO GERAL ══════════ --><p>sample info</p>
<!-- ══════════ SECÇÃO: CHEGADAS RECENTES ══════════ --><p>sample arrival</p>
<!-- ══════════ SECÇÃO: PRODUTOS EM DESTAQUE ══════════ --><p>sample product</p>
<!-- ══════════ SECÇÃO: EM TRÂNSITO ══════════ --><p>sample transit</p>
<!-- ══════════ SECÇÃO: TEMPOS DE ENTREGA ══════════ --><p>delivery</p>
<p style="margin:10px 0 0;font-family:Arial;">Recebe este boletim por ser cliente ou parceiro Vetify.<br><a href="https://vetify.co.ao">Cancelar subscrição</a></p>
</body></html>"""


def make_xlsx_template(path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["D2"] = "old"
    sheet["G3"] = "TOTAL ENCOMENDA"
    sheet["H3"] = "=SUM(H6:H999)"
    headers = [
        "CÓDIGO",
        "DESCRIÇÃO",
        "DISPONIBILIDADE",
        "PVP /S IVA",
        "IVA",
        "PVP C/IVA",
        "QTD A ENC.",
        "VALOR TOTAL",
        "VALIDADE",
        "FAMILIA",
        "SUB-FAMILIA",
        "MARCA / LABORATÓRIO",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(5, column).value = header
    sheet.add_table(Table(displayName="Table1", ref="A5:L10"))
    workbook.create_sheet("Marcas")
    workbook.create_sheet("Dados")
    workbook.save(path)


def test_markdown_renderer_escapes_html_and_allows_only_https_links():
    rendered = markdown_to_email_html(
        "**Safe** <script>alert(1)</script>\n\n- [Good](https://example.com)\n- [Bad](javascript:alert(1))"
    )

    assert "<strong>Safe</strong>" in rendered
    assert "&lt;script&gt;" in rendered
    assert 'href="https://example.com"' in rendered
    assert "javascript:" not in rendered


def test_html_renders_active_sections_and_removes_empty_or_false_blocks():
    data = bulletin(
        announcements=[
            Announcement(
                title="Important <notice>",
                resume="Summary",
                body="**Body**",
                starts_at=dt.date(2026, 9, 7),
            )
        ],
        recent_arrivals=[
            Cargo(
                supplier="Supplier",
                origin="Portugal",
                date=dt.date(2026, 9, 6),
                nature="Medicamentos veterinários",
            )
        ],
    )

    rendered = render_html(html_template(), data)

    assert "Boletim Bissemanal · N.º 73 · 07/09/2026" in rendered
    assert "Important &lt;notice&gt;" in rendered
    assert "<strong>Body</strong>" in rendered
    assert "Medicamentos veterinários" in rendered
    assert "Product One" in rendered
    assert "sample transit" not in rendered
    assert "Cancelar subscrição" not in rendered
    assert "Comunicação operacional" in rendered


def test_html_rejects_false_cancellation_link_it_cannot_replace():
    template = html_template().replace("margin:10px 0 0", "margin:11px 0 0")

    with pytest.raises(RuntimeError, match="false cancellation"):
        render_html(template, bulletin())


def test_cargo_card_never_mentions_skus():
    rendered = render_html(
        html_template(),
        bulletin(
            products=[product(sku="SECRET-SKU")],
            in_transit=[
                Cargo(
                    supplier="Supplier",
                    origin="France",
                    date=dt.date(2026, 10, 1),
                    nature="Alimentação animal",
                )
            ],
        ),
    )

    assert "Alimentação animal" in rendered
    assert "SECRET-SKU" not in rendered


def test_xlsx_uses_correct_rows_formulas_statuses_and_validations(tmp_path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    make_xlsx_template(template)
    data = bulletin(
        products=[
            product(),
            product(
                sku="SKU-2",
                title="Product Two",
                availability="SOB CONSULTA",
                expiry=None,
                used_minimum_stock_fallback=True,
            ),
        ]
    )

    build_xlsx(template, output, data)
    workbook = load_workbook(output, data_only=False)
    sheet = workbook["Sheet1"]

    assert sheet.freeze_panes == "A6"
    assert sheet.tables["Table1"].ref == "A5:L7"
    assert sheet["D2"].value.date() == dt.date(2026, 9, 7)
    assert sheet["F6"].value == "=D6*(1+E6)"
    assert sheet["H6"].value == '=IF(G6="","",G6*F6)'
    assert sheet["H3"].value == "=SUM(H6:H7)"
    assert sheet["C7"].value == "SOB CONSULTA"
    assert sheet["I7"].value is None
    assert sheet["C7"].comment is not None
    assert str(sheet.data_validations.dataValidation[0].sqref) == "G6:G7"
    assert workbook["Marcas"].sheet_state == "hidden"
    assert workbook["Dados"]["A5"].value == "SOB CONSULTA"


def test_email_requires_attachment_and_hides_bulk_recipients(tmp_path):
    attachment = tmp_path / "order.xlsx"
    attachment.write_bytes(b"xlsx")

    message = build_email_message(
        ["one@example.com", "two@example.com"],
        "Subject",
        "<p>Body</p>",
        attachment,
    )

    assert message["To"] is None
    assert message["Bcc"] == "one@example.com, two@example.com"
    assert (
        next(part for part in message.walk() if part.get_filename()).get_filename()
        == "order.xlsx"
    )

    with pytest.raises(FileNotFoundError):
        build_email_message(
            ["one@example.com"], "Subject", "Body", tmp_path / "missing.xlsx"
        )

    with pytest.raises(ValueError, match="invalid email"):
        build_email_message(["bad address"], "Subject", "Body", attachment)
