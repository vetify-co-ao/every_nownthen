import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

os.environ.setdefault("VENDUS_API_KEY", "test-vendus-key")
os.environ.setdefault("SERVICE_ACCOUNT_KEY_PATH", "/tmp/test-service-account.json")
os.environ.setdefault("DASHY_API_KEY", "test-dashy-key")
os.environ["IXLSX_API_KEY"] = ""

import ixlsx


class FakeDashyClient:
    def __init__(self, expiry_by_sku):
        self.expiry_by_sku = expiry_by_sku
        self.requested_skus = []

    def get_nearest_available_lot_expiry(self, sku):
        self.requested_skus.append(sku)
        value = self.expiry_by_sku.get(sku, "")
        if isinstance(value, Exception):
            raise value
        return value


class DashyClientTests(unittest.TestCase):
    def test_stock_lot_request_uses_configured_horizon_and_max_limit(self):
        session = RecordingSession()
        client = ixlsx.DashyClient("test-dashy-key", "https://dashy.example", session)

        expiry = client.get_nearest_available_lot_expiry("SKU-1")

        self.assertEqual(expiry, "2026-08-31")
        self.assertEqual(session.calls[0][0], "https://dashy.example/api/dimensions/product")
        self.assertEqual(session.calls[0][1], {"api_key": "test-dashy-key"})
        self.assertEqual(session.calls[1][0], "https://dashy.example/api/datasets/stock_lots_at_risk")
        self.assertEqual(
            session.calls[1][1],
            {
                "product_sku": "SKU-1",
                "horizon_days": 720,
                "limit": 200,
            },
        )


class RecordingResponse:
    def __init__(self, payload):
        self.payload = payload

    def raise_for_status(self):
        return None

    def json(self):
        return self.payload


class RecordingSession:
    def __init__(self):
        self.headers = {}
        self.calls = []

    def get(self, url, params=None, timeout=None):
        self.calls.append((url, params, timeout))
        if url.endswith("/api/dimensions/product"):
            return RecordingResponse({"rows": []})
        return RecordingResponse(
            {
                "x": ["SKU-1 | Supplier | LOT-A | 2026-08-31 | 2026-01-15"],
                "series": {"remaining_qty": [2], "days_to_expiry": [100]},
            }
        )


class DueDateSelectionTests(unittest.TestCase):
    def test_selects_nearest_available_expiry_from_stock_lots_payload(self):
        payload = {
            "x": [
                "SKU-1 | Supplier | LOT-A | 2026-08-31 | 2026-01-15",
                "SKU-1 | Supplier | LOT-B | 2026-07-15 | 2026-01-20",
                "SKU-1 | Supplier | LOT-C | 2026-06-30 | 2026-01-22",
            ],
            "series": {"remaining_qty": [2, 5, 0], "days_to_expiry": [100, 53, 38]},
        }

        self.assertEqual(ixlsx.select_nearest_lot_expiry(payload), "2026-07-15")

    def test_selects_blank_when_no_available_lot_exists(self):
        payload = {
            "x": ["SKU-1 | Supplier | LOT-A | 2026-08-31 | 2026-01-15"],
            "series": {"remaining_qty": [0], "days_to_expiry": [100]},
        }

        self.assertEqual(ixlsx.select_nearest_lot_expiry(payload), "")


class BuildXlsxDueDateTests(unittest.TestCase):
    def test_stocked_rows_get_nearest_expiry_and_out_of_stock_rows_stay_blank(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            template_path = Path(tmpdir) / "template.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = ixlsx.EXCEL_SHEET_NAME
            sheet["A5"] = "SKU-STOCKED"
            sheet["A6"] = "SKU-EMPTY"
            sheet["I5"] = "old-date"
            sheet["I6"] = "old-date"
            workbook.save(template_path)

            fake_dashy = FakeDashyClient({"SKU-STOCKED": "2026-08-31"})

            with patch.object(ixlsx, "XLSX_TEMPLATE_PATH", str(template_path)), patch.object(
                ixlsx, "OUTPUT_DIR", tmpdir
            ):
                output_path = ixlsx.build_xlsx_file(
                    [
                        {"ProductId": "SKU-STOCKED", "Qty": 3, "NetPrice": 10.0},
                        {"ProductId": "SKU-EMPTY", "Qty": 0, "NetPrice": 20.0},
                    ],
                    dashy_client=fake_dashy,
                )

            output_workbook = load_workbook(output_path)
            output_sheet = output_workbook[ixlsx.EXCEL_SHEET_NAME]

        self.assertEqual(output_sheet["I5"].value, "2026-08-31")
        self.assertIsNone(output_sheet["I6"].value)
        self.assertEqual(fake_dashy.requested_skus, ["SKU-STOCKED"])

    def test_dashy_failures_write_blank_for_stocked_rows(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            template_path = Path(tmpdir) / "template.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = ixlsx.EXCEL_SHEET_NAME
            sheet["A5"] = "SKU-STOCKED"
            sheet["I5"] = "old-date"
            workbook.save(template_path)

            fake_dashy = FakeDashyClient({"SKU-STOCKED": RuntimeError("Dashy down")})

            with patch.object(ixlsx, "XLSX_TEMPLATE_PATH", str(template_path)), patch.object(
                ixlsx, "OUTPUT_DIR", tmpdir
            ):
                output_path = ixlsx.build_xlsx_file(
                    [{"ProductId": "SKU-STOCKED", "Qty": 3, "NetPrice": 10.0}],
                    dashy_client=fake_dashy,
                )

            output_workbook = load_workbook(output_path)
            output_sheet = output_workbook[ixlsx.EXCEL_SHEET_NAME]

        self.assertIsNone(output_sheet["I5"].value)


if __name__ == "__main__":
    unittest.main()
