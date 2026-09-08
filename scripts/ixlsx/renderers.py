from __future__ import annotations

import html
import re
from copy import copy
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from models import Announcement, BulletinData, Cargo, Product

SHEET_NAME = "Sheet1"
FIRST_DATA_ROW = 6
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
HTML_MARKERS = {
    "information": "<!-- ══════════ SECÇÃO: INFORMAÇÃO GERAL ══════════ -->",
    "arrivals": "<!-- ══════════ SECÇÃO: CHEGADAS RECENTES ══════════ -->",
    "highlights": "<!-- ══════════ SECÇÃO: PRODUTOS EM DESTAQUE ══════════ -->",
    "transit": "<!-- ══════════ SECÇÃO: EM TRÂNSITO ══════════ -->",
    "delivery": "<!-- ══════════ SECÇÃO: TEMPOS DE ENTREGA ══════════ -->",
}


def _section_header(title: str, top_padding: int = 20) -> str:
    return f"""<tr>
<td class="vf-pad" style="padding:{top_padding}px 24px 12px;">
<table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0"><tbody><tr>
<td style="white-space:nowrap;"><p style="margin:0;font-family:Arial,Helvetica,sans-serif;font-size:12px;line-height:16px;font-weight:700;text-transform:uppercase;letter-spacing:0.12em;color:#123240;">{html.escape(title)}</p></td>
<td width="100%" style="padding-left:12px;"><table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0"><tbody><tr>
<td width="56" style="border-top:2px solid #123240;font-size:1px;line-height:1px;">&nbsp;</td>
<td width="20" style="border-top:2px solid #33a6d4;font-size:1px;line-height:1px;">&nbsp;</td>
<td style="border-top:1px solid #d5e2e9;font-size:1px;line-height:1px;">&nbsp;</td>
</tr></tbody></table></td>
</tr></tbody></table>
</td>
</tr>"""


def _inline_markdown(source: str) -> str:
    pattern = re.compile(r"\*\*(.+?)\*\*|\[([^\]]+)\]\(([^)]+)\)")
    output = []
    cursor = 0
    for match in pattern.finditer(source):
        output.append(html.escape(source[cursor : match.start()]))
        if match.group(1) is not None:
            output.append(f"<strong>{html.escape(match.group(1))}</strong>")
        else:
            label = html.escape(match.group(2) or "")
            url = match.group(3) or ""
            if url.startswith("https://"):
                output.append(
                    f'<a href="{html.escape(url, quote=True)}" target="_blank" style="color:#1c7fa8;text-decoration:underline;">{label}</a>'
                )
            else:
                output.append(label)
        cursor = match.end()
    output.append(html.escape(source[cursor:]))
    return "".join(output)


def markdown_to_email_html(source: str) -> str:
    blocks = []
    for raw_block in re.split(r"\n\s*\n", source.strip()):
        lines = [line.strip() for line in raw_block.splitlines() if line.strip()]
        if not lines:
            continue
        if all(line.startswith(("- ", "* ")) for line in lines):
            items = "".join(f"<li>{_inline_markdown(line[2:])}</li>" for line in lines)
            blocks.append(
                f'<ul style="margin:0 0 12px;padding-left:22px;font-family:Arial,Helvetica,sans-serif;font-size:15px;line-height:24px;color:#3d5560;">{items}</ul>'
            )
        elif all(re.match(r"^\d+\. ", line) for line in lines):
            items = "".join(
                "<li>" + _inline_markdown(re.sub(r"^\d+\. ", "", line)) + "</li>"
                for line in lines
            )
            blocks.append(
                f'<ol style="margin:0 0 12px;padding-left:22px;font-family:Arial,Helvetica,sans-serif;font-size:15px;line-height:24px;color:#3d5560;">{items}</ol>'
            )
        elif len(lines) == 1 and re.match(r"^#{1,3}\s+", lines[0]):
            heading = re.sub(r"^#{1,3}\s+", "", lines[0])
            blocks.append(
                f'<h3 style="margin:0 0 10px;font-family:Arial,Helvetica,sans-serif;font-size:17px;line-height:23px;color:#123240;">{_inline_markdown(heading)}</h3>'
            )
        else:
            content = "<br>".join(_inline_markdown(line) for line in lines)
            blocks.append(
                f'<p style="margin:0 0 12px;font-family:Arial,Helvetica,sans-serif;font-size:15px;line-height:24px;color:#3d5560;">{content}</p>'
            )
    return "".join(blocks)


def _announcement_card(item: Announcement) -> str:
    resume = (
        f'<p style="margin:0 0 12px;font-family:Arial,Helvetica,sans-serif;font-size:15px;line-height:24px;font-weight:700;color:#3d5560;">{html.escape(item.resume)}</p>'
        if item.resume
        else ""
    )
    return f"""<tr>
<td class="vf-pad" style="padding:0 24px 14px;">
<table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0" bgcolor="#FFFFFF" style="border-collapse:separate;border:1px solid #d5e2e9;border-radius:18px;overflow:hidden;background-color:#ffffff;width:100%;"><tbody><tr><td style="padding:24px;">
<table role="presentation" cellspacing="0" cellpadding="0" border="0" style="margin-bottom:12px;"><tbody><tr><td><span style="display:inline-block;padding:4px 10px;font-family:Arial,Helvetica,sans-serif;font-size:11px;line-height:16px;font-weight:700;text-transform:uppercase;letter-spacing:0.1em;border-radius:9999px;color:#1c7fa8;background-color:#e9f5fa;border:1px solid #bfe2f0;">Comunicado</span></td></tr></tbody></table>
<h2 style="margin:0 0 12px;font-family:Arial,Helvetica,sans-serif;font-size:21px;line-height:27px;mso-line-height-rule:exactly;font-weight:700;letter-spacing:-0.01em;color:#123240;">{html.escape(item.title)}</h2>
{resume}{markdown_to_email_html(item.body)}
</td></tr></tbody></table>
</td>
</tr>"""


def _product_card(product: Product) -> str:
    title = html.escape(product.title)
    image = ""
    if product.image_url:
        image = f'''<tr><td bgcolor="#FFFFFF" style="border-bottom:1px solid #e8f0f4;background-color:#ffffff;background-image:linear-gradient(#ffffff,#ffffff);" align="center"><img src="{html.escape(product.image_url, quote=True)}" alt="{title}" width="400" style="width:100%;max-width:400px;height:auto;display:block;background-color:#ffffff;"></td></tr>'''
    description = (
        f'<p style="margin:0 0 16px;font-family:Arial,Helvetica,sans-serif;font-size:15px;line-height:24px;color:#3d5560;">{html.escape(product.description).replace(chr(10), "<br>")}</p>'
        if product.description
        else ""
    )
    cta = ""
    if product.slug:
        url = f"https://www.vetify.co.ao/products/{quote(product.slug, safe='-._~')}"
        cta = f'''<table role="presentation" cellspacing="0" cellpadding="0" border="0"><tbody><tr><td bgcolor="#33a6d4" style="border-radius:9999px;"><a href="{html.escape(url, quote=True)}" target="_blank" style="display:block;padding:10px 22px;font-family:Arial,Helvetica,sans-serif;font-size:14px;line-height:20px;font-weight:700;color:#ffffff;text-decoration:none;border-radius:9999px;">Saiba mais&nbsp;→</a></td></tr></tbody></table>'''
    return f"""<tr>
<td class="vf-pad" style="padding:0 24px 14px;">
<table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0" bgcolor="#FFFFFF" style="border-collapse:separate;border:1px solid #d5e2e9;border-radius:18px;overflow:hidden;background-color:#ffffff;width:100%;"><tbody>
{image}<tr><td style="padding:24px;">
<table role="presentation" cellspacing="0" cellpadding="0" border="0" style="margin-bottom:12px;"><tbody><tr>
<td style="padding-right:8px;"><span style="display:inline-block;padding:4px 10px;font-family:Arial,Helvetica,sans-serif;font-size:11px;line-height:16px;font-weight:700;text-transform:uppercase;letter-spacing:0.1em;border-radius:4px;color:#ffffff;background-color:#123240;">Destaque</span></td>
<td><span style="display:inline-block;padding:4px 10px;font-family:Arial,Helvetica,sans-serif;font-size:11px;line-height:16px;font-weight:700;text-transform:uppercase;letter-spacing:0.1em;border-radius:9999px;color:#1c7fa8;background-color:#e9f5fa;border:1px solid #bfe2f0;">{html.escape(product.department)}</span></td>
</tr></tbody></table>
<h2 style="margin:0 0 12px;font-family:Arial,Helvetica,sans-serif;font-size:20px;line-height:26px;mso-line-height-rule:exactly;font-weight:700;letter-spacing:-0.01em;color:#123240;">{title}</h2>
{description}{cta}
</td></tr></tbody></table>
</td>
</tr>"""


def _cargo_card(cargo: Cargo, recent: bool) -> str:
    badge = "Em armazém" if recent else "Em trânsito"
    badge_colors = (
        ("#1e7d4f", "#e8f6ee", "#bfe6d0")
        if recent
        else ("#8a5a12", "#fdf3e2", "#f0d9ae")
    )
    rows = []
    if cargo.origin:
        rows.append(("Origem", cargo.origin))
    if cargo.date:
        rows.append(
            (
                "Desalfandegado em" if recent else "Chegada prevista",
                cargo.date.strftime("%d/%m/%Y"),
            )
        )
    if cargo.nature:
        rows.append(("Natureza da carga", cargo.nature))
    rows_html = "".join(
        f"""<tr><td width="140" style="padding:7px 0;border-top:1px solid #e8f0f4;font-family:Arial,Helvetica,sans-serif;font-size:12px;line-height:18px;font-weight:700;text-transform:uppercase;letter-spacing:0.06em;color:#5b7683;vertical-align:top;">{html.escape(label)}</td><td style="padding:7px 0;border-top:1px solid #e8f0f4;font-family:Arial,Helvetica,sans-serif;font-size:14px;line-height:21px;color:#123240;">{html.escape(value)}</td></tr>"""
        for label, value in rows
    )
    return f"""<tr><td class="vf-pad" style="padding:0 24px 14px;">
<table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0" bgcolor="#FFFFFF" style="border-collapse:separate;border:1px solid #d5e2e9;border-radius:18px;overflow:hidden;background-color:#ffffff;width:100%;"><tbody>
<tr><td style="padding:20px 24px 8px;"><table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0"><tbody><tr>
<td><h3 style="margin:0;font-family:Arial,Helvetica,sans-serif;font-size:17px;line-height:23px;font-weight:700;color:#123240;">{html.escape(cargo.supplier)}</h3></td>
<td align="right"><span style="display:inline-block;padding:4px 10px;font-family:Arial,Helvetica,sans-serif;font-size:11px;line-height:16px;font-weight:700;text-transform:uppercase;letter-spacing:0.1em;border-radius:9999px;color:{badge_colors[0]};background-color:{badge_colors[1]};border:1px solid {badge_colors[2]};">{badge}</span></td>
</tr></tbody></table></td></tr>
<tr><td style="padding:8px 24px 20px;"><table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0"><tbody>{rows_html}</tbody></table></td></tr>
</tbody></table></td></tr>"""


def _replace_section(document: str, start: str, end: str, replacement: str) -> str:
    pattern = re.compile(
        re.escape(start) + r".*?(?=" + re.escape(end) + r")", re.DOTALL
    )
    updated, count = pattern.subn(
        replacement.rstrip() + "\n\n" if replacement else "", document, count=1
    )
    if count != 1:
        raise RuntimeError(
            f"Required HTML template marker is missing or duplicated: {start}"
        )
    return updated


def render_html(template_html: str, bulletin: BulletinData) -> str:
    if not template_html.strip():
        raise RuntimeError("VCRM HTML template is empty")
    for marker in HTML_MARKERS.values():
        if template_html.count(marker) != 1:
            raise RuntimeError(
                f"Required HTML template marker is missing or duplicated: {marker}"
            )

    date_display = bulletin.publication_date.strftime("%d/%m/%Y")
    document = re.sub(
        r"<title>.*?</title>",
        "<title>Boletim Bissemanal Vetify</title>",
        template_html,
        count=1,
        flags=re.DOTALL,
    )
    document, count = re.subn(
        r"Boletim Semanal\s*·\s*N\.º\s*00\s*·\s*00/00/0000",
        f"Boletim Bissemanal · N.º {bulletin.issue_number} · {date_display}",
        document,
        count=1,
    )
    if count != 1:
        raise RuntimeError("VCRM HTML template bulletin header placeholder is missing")

    information = ""
    if bulletin.announcements:
        information = (
            HTML_MARKERS["information"]
            + "\n"
            + _section_header("Informação geral", 28)
            + "\n"
            + "\n".join(_announcement_card(item) for item in bulletin.announcements)
        )
    arrivals = ""
    if bulletin.recent_arrivals:
        arrivals = (
            HTML_MARKERS["arrivals"]
            + "\n"
            + _section_header("Chegadas recentes")
            + "\n"
            + "\n".join(_cargo_card(item, True) for item in bulletin.recent_arrivals)
        )
    highlights = ""
    featured = [product for product in bulletin.products if product.featured]
    if featured:
        highlights = (
            HTML_MARKERS["highlights"]
            + "\n"
            + _section_header("Produtos em destaque")
            + "\n"
            + "\n".join(_product_card(item) for item in featured)
        )
    transit = ""
    if bulletin.in_transit:
        transit = (
            HTML_MARKERS["transit"]
            + "\n"
            + _section_header("Em trânsito")
            + "\n"
            + "\n".join(_cargo_card(item, False) for item in bulletin.in_transit)
        )

    document = _replace_section(
        document, HTML_MARKERS["information"], HTML_MARKERS["arrivals"], information
    )
    document = _replace_section(
        document, HTML_MARKERS["arrivals"], HTML_MARKERS["highlights"], arrivals
    )
    document = _replace_section(
        document, HTML_MARKERS["highlights"], HTML_MARKERS["transit"], highlights
    )
    document = _replace_section(
        document, HTML_MARKERS["transit"], HTML_MARKERS["delivery"], transit
    )
    document = re.sub(
        r'<p style="margin:10px 0 0;[^>]*">Recebe este boletim por ser cliente ou parceiro Vetify\.<br><a [^>]*>Cancelar subscrição</a></p>',
        '<p style="margin:10px 0 0;font-family:Arial,Helvetica,sans-serif;font-size:11px;line-height:17px;color:#8ba2ac;">Comunicação operacional destinada a clientes e parceiros Vetify.</p>',
        document,
        count=1,
    )
    if "Cancelar subscrição" in document:
        raise RuntimeError(
            "VCRM HTML template still contains a false cancellation link"
        )
    return document


def build_xlsx(template_path: Path, output_path: Path, bulletin: BulletinData) -> Path:
    workbook = load_workbook(template_path)
    required_sheets = {SHEET_NAME, "Marcas", "Dados"}
    if not required_sheets.issubset(workbook.sheetnames):
        raise RuntimeError(
            f"VCRM XLSX template lacks sheets: {sorted(required_sheets - set(workbook.sheetnames))}"
        )
    sheet = workbook[SHEET_NAME]
    if "Table1" not in sheet.tables:
        raise RuntimeError("VCRM XLSX template lacks Table1")
    if not bulletin.products:
        raise RuntimeError("Cannot build an empty product workbook")

    sheet["D2"] = bulletin.publication_date
    sheet["D2"].number_format = "dd/mm/yyyy"
    sheet.freeze_panes = "A6"
    last_row = FIRST_DATA_ROW + len(bulletin.products) - 1
    existing_last_row = sheet.max_row
    if last_row > existing_last_row:
        for row in range(existing_last_row + 1, last_row + 1):
            for column in range(1, 13):
                sheet.cell(row, column)._style = copy(
                    sheet.cell(FIRST_DATA_ROW, column)._style
                )

    for row in range(FIRST_DATA_ROW, max(existing_last_row, last_row) + 1):
        for column in range(1, 13):
            sheet.cell(row, column).value = None

    for row, product in enumerate(bulletin.products, start=FIRST_DATA_ROW):
        values = [
            product.sku,
            product.title,
            product.availability,
            product.net_price,
            product.vat_rate,
            f"=D{row}*(1+E{row})",
            None,
            f'=IF(G{row}="","",G{row}*F{row})',
            product.expiry,
            product.categories,
            product.subcategories,
            product.brand,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column).value = value
        for column in (4, 6, 8):
            sheet.cell(row, column).number_format = '#,##0.00\\ "Kz"'
        sheet.cell(row, 5).number_format = "0%"
        sheet.cell(row, 9).number_format = "yyyy-mm-dd"
        if product.used_minimum_stock_fallback:
            sheet.cell(row, 3).comment = Comment(
                "Stock mínimo VPIM ausente/inválido; disponibilidade calculada com fallback 10.",
                "Vetify",
            )

    sheet["H3"] = f"=SUM(H{FIRST_DATA_ROW}:H{last_row})"
    sheet.tables["Table1"].ref = f"A5:L{last_row}"
    if sheet.max_row > last_row:
        sheet.delete_rows(last_row + 1, sheet.max_row - last_row)

    sheet.data_validations.dataValidation = []
    quantity_validation = DataValidation(
        type="whole",
        operator="between",
        formula1="0",
        formula2="999999",
        allow_blank=True,
    )
    quantity_validation.promptTitle = "Quantidade a encomendar"
    quantity_validation.prompt = "Introduza um número inteiro igual ou superior a zero."
    quantity_validation.errorTitle = "Quantidade inválida"
    quantity_validation.error = "Use um número inteiro entre 0 e 999999."
    quantity_validation.errorStyle = "stop"
    quantity_validation.showErrorMessage = True
    quantity_validation.showInputMessage = True
    sheet.add_data_validation(quantity_validation)
    quantity_validation.add(f"G{FIRST_DATA_ROW}:G{last_row}")

    sheet.conditional_formatting._cf_rules.clear()
    for status, fill_color, font_color in (
        ("EM STOCK", "E8F6EE", "1E7D4F"),
        ("ULTIMAS UNIDADES", "FFF3CD", "8A5A12"),
        ("ESGOTADO", "FFC7CE", "9C0006"),
        ("SOB CONSULTA", "E9F5FA", "1C7FA8"),
    ):
        sheet.conditional_formatting.add(
            f"C{FIRST_DATA_ROW}:C{last_row}",
            FormulaRule(
                formula=[f'$C{FIRST_DATA_ROW}="{status}"'],
                stopIfTrue=True,
                fill=PatternFill("solid", fgColor=fill_color),
                font=Font(color=font_color),
            ),
        )

    sheet["C5"].comment = Comment(
        "Fonte: Vendus. Sem Vendus = SOB CONSULTA; stock <= 0 = ESGOTADO; "
        "stock <= mínimo VPIM = ULTIMAS UNIDADES; acima do mínimo = EM STOCK. "
        "Stock mínimo em falta usa fallback 10.",
        "Vetify",
    )
    sheet["D5"].comment = Comment("Fonte: VPIM product:price.", "Vetify")
    sheet["E5"].comment = Comment("Fonte: VPIM product:vat-rate.", "Vetify")
    sheet["I5"].comment = Comment(
        "Fonte: Dashy; validade futura mais próxima com quantidade remanescente positiva.",
        "Vetify",
    )
    sheet["J5"].comment = Comment(
        "Fonte: taxonomia VPIM; múltiplos valores separados por ·.", "Vetify"
    )

    brands_sheet = workbook["Marcas"]
    brands_sheet.delete_rows(1, brands_sheet.max_row)
    for row, brand in enumerate(
        sorted(
            {product.brand for product in bulletin.products if product.brand},
            key=str.casefold,
        ),
        start=1,
    ):
        brands_sheet.cell(row, 1).value = brand
    brands_sheet.sheet_state = "hidden"

    data_sheet = workbook["Dados"]
    data_sheet.delete_rows(1, data_sheet.max_row)
    for row, status in enumerate(
        ["DISPONIBILIDADE", "EM STOCK", "ULTIMAS UNIDADES", "ESGOTADO", "SOB CONSULTA"],
        start=1,
    ):
        data_sheet.cell(row, 1).value = status
    data_sheet.sheet_state = "hidden"

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def build_email_message(
    recipients: list[str],
    subject: str,
    html_body: str,
    attachment_path: Path,
    simulation: bool = False,
) -> EmailMessage:
    if not attachment_path.is_file():
        raise FileNotFoundError(
            f"Required XLSX attachment not found: {attachment_path}"
        )
    if any(not EMAIL_PATTERN.fullmatch(recipient) for recipient in recipients):
        raise ValueError("Recipient list contains an invalid email address")
    message = EmailMessage()
    message["From"] = "Vetify <comercial@vetify.co.ao>"
    message["Reply-To"] = "encomendas@vetify.co.ao"
    message["Subject"] = subject
    if simulation:
        message["To"] = "undisclosed-recipients:;"
        message["X-Vetify-Simulation"] = "true"
        message["X-Simulated-Bcc-Count"] = str(len(recipients))
    elif recipients:
        message["Bcc"] = ", ".join(recipients)
    else:
        raise ValueError("At least one recipient is required for delivery")
    message.set_content("Esta mensagem requer um cliente de email com suporte HTML.")
    message.add_alternative(html_body, subtype="html")
    message.add_attachment(
        attachment_path.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=attachment_path.name,
    )
    return message
